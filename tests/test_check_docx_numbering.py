import importlib.util
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

SCRIPT = (
    Path(__file__).resolve().parents[1]
    / "no-markdown-bullets"
    / "scripts"
    / "check_docx_numbering.py"
)
SPEC = importlib.util.spec_from_file_location("check_docx_numbering", SCRIPT)
CHECKER = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = CHECKER
SPEC.loader.exec_module(CHECKER)

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def paragraph(text, num_id=None, level="0", style=None):
    properties = ""
    if style is not None:
        properties += f'<w:pStyle w:val="{style}"/>'
    if num_id is not None:
        properties += (
            f'<w:numPr><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'
        )
    ppr = f"<w:pPr>{properties}</w:pPr>" if properties else ""
    return f"<w:p>{ppr}<w:r><w:t>{text}</w:t></w:r></w:p>"


def numbering_xml(abstract_levels, numbers):
    """Build numbering.xml.

    ``abstract_levels`` is {abstractId: {level: numFmt}} and ``numbers`` is
    {numId: abstractId}.
    """
    abstracts = []
    for abstract_id, levels in abstract_levels.items():
        lvls = "".join(
            f'<w:lvl w:ilvl="{level}"><w:numFmt w:val="{fmt}"/></w:lvl>'
            for level, fmt in levels.items()
        )
        abstracts.append(
            f'<w:abstractNum w:abstractNumId="{abstract_id}">{lvls}</w:abstractNum>'
        )
    nums = "".join(
        f'<w:num w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'
        for num_id, abstract_id in numbers.items()
    )
    return (
        f'<?xml version="1.0"?><w:numbering xmlns:w="{W}">'
        + "".join(abstracts)
        + nums
        + "</w:numbering>"
    )


def styles_xml(styles):
    """Build styles.xml from {styleId: (numId or None, basedOn or None)}."""
    body = []
    for style_id, (num_id, based_on) in styles.items():
        inner = ""
        if based_on:
            inner += f'<w:basedOn w:val="{based_on}"/>'
        if num_id:
            inner += (
                f'<w:pPr><w:numPr><w:ilvl w:val="0"/>'
                f'<w:numId w:val="{num_id}"/></w:numPr></w:pPr>'
            )
        body.append(f'<w:style w:styleId="{style_id}">{inner}</w:style>')
    return (
        f'<?xml version="1.0"?><w:styles xmlns:w="{W}">'
        + "".join(body)
        + "</w:styles>"
    )


def build_docx(path, paragraphs, numbering=None, styles=None):
    document = (
        f'<?xml version="1.0"?><w:document xmlns:w="{W}"><w:body>'
        + "".join(paragraphs)
        + "</w:body></w:document>"
    )
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document)
        if numbering is not None:
            archive.writestr("word/numbering.xml", numbering)
        if styles is not None:
            archive.writestr("word/styles.xml", styles)


BULLETS = numbering_xml(
    {"0": {"0": "bullet", "1": "bullet"}, "1": {"0": "decimal"}},
    {"1": "0", "2": "1"},
)


class ScanTests(unittest.TestCase):
    def scan(self, paragraphs, numbering=BULLETS, styles=None):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "doc.docx"
            build_docx(path, paragraphs, numbering, styles)
            return CHECKER.scan_docx(path)

    def test_direct_bullet_is_reported(self):
        result = self.scan([paragraph("first", num_id="1")])
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].index, 1)
        self.assertEqual(result.findings[0].text, "first")
        self.assertEqual(result.findings[0].source, "direct numPr")

    def test_decimal_numbering_is_not_reported(self):
        result = self.scan([paragraph("step", num_id="2")])
        self.assertEqual(result.findings, [])
        self.assertEqual(result.numbered, 1)

    def test_unnumbered_paragraphs_are_ignored(self):
        result = self.scan([paragraph("prose only")])
        self.assertEqual(result.findings, [])
        self.assertEqual(result.paragraphs, 1)
        self.assertEqual(result.numbered, 0)

    def test_num_id_zero_means_no_numbering(self):
        result = self.scan([paragraph("plain", num_id="0")])
        self.assertEqual(result.findings, [])

    def test_level_selects_the_matching_format(self):
        result = self.scan([paragraph("nested", num_id="1", level="1")])
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].level, "1")

    def test_style_numbering_is_resolved(self):
        styles = styles_xml({"ListBullet": ("1", None)})
        result = self.scan([paragraph("styled", style="ListBullet")], styles=styles)
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].source, "style ListBullet")

    def test_based_on_chain_is_followed(self):
        styles = styles_xml({
            "CustomBullet": (None, "ListBullet"),
            "ListBullet": ("1", None),
        })
        result = self.scan(
            [paragraph("inherited", style="CustomBullet")], styles=styles
        )
        self.assertEqual(len(result.findings), 1)
        self.assertEqual(result.findings[0].source, "style CustomBullet")

    def test_direct_numbering_wins_over_style(self):
        styles = styles_xml({"ListBullet": ("1", None)})
        result = self.scan(
            [paragraph("decimal wins", num_id="2", style="ListBullet")], styles=styles
        )
        self.assertEqual(result.findings, [])

    def test_based_on_cycle_terminates(self):
        styles = styles_xml({"A": (None, "B"), "B": (None, "A")})
        result = self.scan([paragraph("loop", style="A")], styles=styles)
        self.assertEqual(result.findings, [])

    def test_unknown_style_is_ignored(self):
        result = self.scan([paragraph("ghost", style="Missing")], styles=styles_xml({}))
        self.assertEqual(result.findings, [])

    def test_missing_numbering_part_yields_no_findings(self):
        result = self.scan([paragraph("plain", num_id="1")], numbering=None)
        self.assertEqual(result.findings, [])

    def test_long_text_is_truncated(self):
        result = self.scan([paragraph("x" * 200, num_id="1")])
        self.assertTrue(result.findings[0].text.endswith("..."))
        self.assertLessEqual(len(result.findings[0].text), CHECKER.SNIPPET_LIMIT + 3)

    def test_whitespace_in_text_is_collapsed(self):
        result = self.scan([paragraph("two    spaces", num_id="1")])
        self.assertEqual(result.findings[0].text, "two spaces")


class CommandTests(unittest.TestCase):
    def test_clean_file_exits_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "clean.docx"
            build_docx(
                path, [paragraph("prose"), paragraph("step", num_id="2")], BULLETS
            )
            self.assertEqual(CHECKER.run([str(path)], stdout=io.StringIO()), 0)

    def test_bullet_file_exits_one(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bullets.docx"
            build_docx(path, [paragraph("item", num_id="1")], BULLETS)
            output = io.StringIO()
            self.assertEqual(CHECKER.run([str(path)], stdout=output), 1)
            self.assertIn("paragraph 1: bullet numbering", output.getvalue())

    def test_missing_path_exits_two(self):
        self.assertEqual(
            CHECKER.run(
                ["absent.docx"], stdout=io.StringIO(), stderr=io.StringIO()
            ),
            CHECKER.EXIT_ERROR,
        )

    def test_not_a_zip_exits_two(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "fake.docx"
            path.write_text("not a zip", encoding="utf-8")
            stderr = io.StringIO()
            self.assertEqual(
                CHECKER.run([str(path)], stdout=io.StringIO(), stderr=stderr),
                CHECKER.EXIT_ERROR,
            )
            self.assertIn("fake.docx", stderr.getvalue())

    def test_missing_document_part_exits_two(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "empty.docx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("word/styles.xml", styles_xml({}))
            self.assertEqual(
                CHECKER.run([str(path)], stdout=io.StringIO(), stderr=io.StringIO()),
                CHECKER.EXIT_ERROR,
            )

    def test_directory_scan_finds_nested_files(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "sub").mkdir()
            build_docx(root / "a.docx", [paragraph("one", num_id="1")], BULLETS)
            build_docx(root / "sub" / "b.docx", [paragraph("two", num_id="1")], BULLETS)
            output = io.StringIO()
            self.assertEqual(CHECKER.run([str(root)], stdout=output), 1)
            self.assertEqual(len(output.getvalue().strip().splitlines()), 2)

    def test_json_output_shape(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bullets.docx"
            build_docx(
                path,
                [paragraph("item", num_id="1"), paragraph("step", num_id="2")],
                BULLETS,
            )
            output = io.StringIO()
            status = CHECKER.run([str(path)], output_format="json", stdout=output)
            self.assertEqual(status, 1)
            payload = json.loads(output.getvalue())
            self.assertEqual(payload["summary"]["findings"], 1)
            self.assertEqual(payload["summary"]["files"], 1)
            record = payload["files"][0]
            self.assertEqual(record["numbered"], 2)
            self.assertEqual(record["findings"][0]["num_id"], "1")


try:
    import docx  # noqa: F401
except ImportError:  # pragma: no cover - optional test dependency
    docx = None


@unittest.skipIf(docx is None, "python-docx is not installed")
class RealDocumentTests(unittest.TestCase):
    """Check a document written by python-docx, not by our own test helper.

    The hand-built fixtures prove the parser handles the shapes we thought of.
    This proves it handles the shapes a real library emits, including the
    default template's "List Bullet" and "List Number" styles.
    """

    def build(self, path, bullet_items, number_items, extra_prose=True):
        document = docx.Document()
        if extra_prose:
            document.add_paragraph("Ordinary prose paragraph.")
        for text in bullet_items:
            document.add_paragraph(text, style="List Bullet")
        for text in number_items:
            document.add_paragraph(text, style="List Number")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Field"
        table.cell(1, 0).text = "alpha"
        document.save(str(path))

    def test_default_template_bullets_are_found(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "real.docx"
            self.build(path, ["Fix timeout", "Reduce startup"], ["Step one"])
            result = CHECKER.scan_docx(path)
            self.assertEqual(len(result.findings), 2)
            self.assertTrue(
                all("ListBullet" in finding.source for finding in result.findings)
            )
            self.assertEqual(result.findings[0].text, "Fix timeout")

    def test_default_template_numbering_is_not_reported(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "numbers.docx"
            self.build(path, [], ["Step one", "Step two"])
            result = CHECKER.scan_docx(path)
            self.assertEqual(result.findings, [])
            self.assertEqual(result.numbered, 2)

    def test_clean_document_exits_zero(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "clean.docx"
            self.build(path, [], ["Only ordered"])
            self.assertEqual(CHECKER.run([str(path)], stdout=io.StringIO()), 0)


if __name__ == "__main__":
    unittest.main()
